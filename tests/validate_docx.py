"""Validates the downloaded Pricing_Decision_Report.docx and checks it against Pricing_Decision_Data.json.

Usage: python tests/validate_docx.py <report.docx> <data.json>
Prints DOCX_VALID and CONSISTENT when all checks pass.
"""
import json
import sys
import zipfile

import docx  # python-docx

report, data_path = sys.argv[1], sys.argv[2]
ok = True


def check(cond, label):
    global ok
    print(("ok   " if cond else "FAIL ") + label)
    ok = ok and cond


with zipfile.ZipFile(report) as z:
    names = z.namelist()
    check(z.testzip() is None, "zip container intact")
    check("[Content_Types].xml" in names and "word/document.xml" in names, "Office Open XML parts present")

d = docx.Document(report)
# Every style referenced by the document must be defined (unstyled paragraphs use Word's built-in Normal).
import re
with zipfile.ZipFile(report) as z:
    doc_xml = z.read("word/document.xml").decode("utf-8")
    styles_xml = z.read("word/styles.xml").decode("utf-8")
used = set(re.findall(r'w:(?:pStyle|rStyle) w:val="([^"]+)"', doc_xml))
defined = set(re.findall(r'w:styleId="([^"]+)"', styles_xml))
missing = sorted(used - defined)
check(not missing, f"all referenced styles are defined ({sorted(used)}; missing: {missing})")
headings = [p.text for p in d.paragraphs if p.style is not None and p.style.name == "Heading 1"]
expected = ["1. Venture Overview", "2. Website Information", "3. Final Pricing Decision", "4. Cost Assumptions", "5. Financial Results",
            "6. Competitive Positioning", "7. Pricing Decision Review", "8. Student Decision Statement", "9. Sources and Limitations", "10. Confirmation Record"]
check(headings == expected, f"10 report sections in order ({len(headings)} found)")
check(len(d.tables) >= 12, f"tables present ({len(d.tables)})")
text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
with open(data_path, encoding="utf-8") as f:
    data = json.load(f)
price = data["finalPrices"][0]["price"]
check(f"${price:,.2f}" in text, f"final price ${price:,.2f} appears in the report")
check(data["website"]["repositoryUrl"] in text and data["website"]["liveUrl"] in text, "both website URLs in the report")
check(data["studentJustification"] in text, "student justification reproduced")
check(data["riskFindings"]["classification"] in text, "risk classification matches JSON")
check(data["confirmation"]["snapshotFingerprint"] in text, "same snapshot ID as JSON")
profit = data["financialAssumptions"]["results"]["operatingProfit"]
check(f"${profit:,.2f}".replace("$-", "-$") in text, f"operating profit {profit:,.2f} matches JSON")
active = [e["name"] for e in data["financialAssumptions"]["operatingExpenses"]]
check(all(n in text for n in active), f"all {len(active)} active operating expenses listed")
startup = [e["name"] for e in data["financialAssumptions"]["startupCosts"]]
check(all(n in text for n in startup), f"all {len(startup)} startup costs listed")
print("DOCX_VALID" if ok else "DOCX_INVALID")
print("CONSISTENT" if ok else "INCONSISTENT")
